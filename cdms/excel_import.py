from io import BytesIO

from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from cdms.serializers import ProductSnapshotSerializer
from cdms.service import CdmsService


class ExcelProductImporter:
    """Read Product rows from an .xlsx file and send each row to CDMS."""

    def __init__(self):
        self.cdms_service = CdmsService()

    def import_file(self, uploaded_file) -> dict:
        if not uploaded_file.name.lower().endswith(".xlsx"):
            raise ValidationError({"file": "Only .xlsx files are supported."})

        workbook = load_workbook(
            BytesIO(uploaded_file.read()),
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            raise ValidationError({"file": "The Excel file has no header row."})

        headers = [str(header).strip() if header is not None else "" for header in headers]
        changed = 0
        unchanged = 0
        errors = []

        for row_number, values in enumerate(rows, start=2):
            if not any(value is not None for value in values):
                continue

            product_data = {
                header: value
                for header, value in zip(headers, values)
                if header
            }
            serializer = ProductSnapshotSerializer(data=product_data)

            if not serializer.is_valid():
                errors.append({"row": row_number, "errors": serializer.errors})
                continue

            _, row_changed = self.cdms_service.sync_product(
                serializer.validated_data
            )
            if row_changed:
                changed += 1
            else:
                unchanged += 1

        workbook.close()
        return {
            "changed": changed,
            "unchanged": unchanged,
            "errors": errors,
        }
